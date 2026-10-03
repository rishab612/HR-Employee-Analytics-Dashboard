import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================================
# FILES
# ==========================================================

input_file = "HR_Employee_Data.xlsx"
output_file = "HR_Employee_Analytics.xlsx"

# ==========================================================
# READ EMPLOYEE DATA
# ==========================================================

df = pd.read_excel(input_file)

# ==========================================================
# DEPARTMENT TABLE
# ==========================================================

departments = sorted(df["Department"].unique())

department_df = pd.DataFrame({
    "DepartmentID": [
        f"DEPT{i+1:03d}"
        for i in range(len(departments))
    ],
    "DepartmentName": departments
})

# ==========================================================
# JOB ROLE TABLE
# ==========================================================

roles = sorted(df["JobRole"].unique())

role_df = pd.DataFrame({
    "JobRoleID": [
        f"ROLE{i+1:03d}"
        for i in range(len(roles))
    ],
    "JobRole": roles
})

# ==========================================================
# DATE TABLE
# ==========================================================

start_date = pd.to_datetime(df["HireDate"]).min()
end_date = pd.Timestamp("2026-12-31")

dates = pd.date_range(
    start=start_date,
    end=end_date,
    freq="D"
)

date_df = pd.DataFrame({
    "Date": dates
})

date_df["Year"] = date_df["Date"].dt.year
date_df["Month Number"] = date_df["Date"].dt.month
date_df["Month"] = date_df["Date"].dt.strftime("%B")
date_df["Quarter"] = (
    "Q" +
    date_df["Date"].dt.quarter.astype(str)
)
date_df["Year Month"] = (
    date_df["Date"].dt.strftime("%Y-%m")
)

# ==========================================================
# DATA DICTIONARY
# ==========================================================

data_dictionary = [

    ["EmployeeID", "Unique employee identifier", "Text"],
    ["EmployeeName", "Employee full name", "Text"],
    ["Gender", "Employee gender", "Category"],
    ["Age", "Employee age", "Number"],
    ["Department", "Employee department", "Category"],
    ["JobRole", "Employee job role", "Category"],
    ["Education", "Highest education level", "Category"],
    ["JobLevel", "Employee organizational level", "Number"],
    ["MonthlyIncome", "Monthly employee salary", "Currency"],
    ["TotalWorkingYears", "Total professional experience", "Number"],
    ["YearsAtCompany", "Years employed at company", "Number"],
    ["YearsInCurrentRole", "Years in current role", "Number"],
    ["YearsSincePromotion", "Years since last promotion", "Number"],
    ["JobSatisfaction", "Job satisfaction score 1-4", "Number"],
    ["EnvironmentSatisfaction", "Work environment satisfaction 1-4", "Number"],
    ["WorkLifeBalance", "Work-life balance score 1-4", "Number"],
    ["PerformanceRating", "Employee performance rating", "Number"],
    ["TrainingHours", "Annual training hours", "Number"],
    ["Overtime", "Whether employee works overtime", "Category"],
    ["BusinessTravel", "Business travel frequency", "Category"],
    ["MaritalStatus", "Marital status", "Category"],
    ["DistanceFromHome", "Distance from home in km", "Number"],
    ["NumCompaniesWorked", "Number of previous employers", "Number"],
    ["HireDate", "Employee joining date", "Date"],
    ["ExitDate", "Employee exit date if applicable", "Date"],
    ["Attrition", "Whether employee left the company", "Category"]
]

dictionary_df = pd.DataFrame(
    data_dictionary,
    columns=[
        "Column Name",
        "Description",
        "Data Type"
    ]
)

# ==========================================================
# WRITE EXCEL FILE
# ==========================================================

with pd.ExcelWriter(
    output_file,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Employee Data",
        index=False
    )

    department_df.to_excel(
        writer,
        sheet_name="Department",
        index=False
    )

    role_df.to_excel(
        writer,
        sheet_name="Job Roles",
        index=False
    )

    date_df.to_excel(
        writer,
        sheet_name="Date Table",
        index=False
    )

    dictionary_df.to_excel(
        writer,
        sheet_name="Data Dictionary",
        index=False
    )

# ==========================================================
# FORMAT WORKBOOK
# ==========================================================

wb = load_workbook(output_file)

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

thin_border = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)

for ws in wb.worksheets:

    # Freeze first row
    ws.freeze_panes = "A2"

    # Header formatting
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Column widths
    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            35
        )

    # Add borders
    for row in ws.iter_rows():

        for cell in row:

            cell.border = thin_border

# ==========================================================
# NUMBER FORMATTING
# ==========================================================

ws = wb["Employee Data"]

# Find columns
headers = {
    cell.value: cell.column
    for cell in ws[1]
}

# Salary formatting
if "MonthlyIncome" in headers:

    col = headers["MonthlyIncome"]

    for row in range(2, ws.max_row + 1):

        ws.cell(row, col).number_format = '₹#,##0'

# Date formatting
for date_column in ["HireDate", "ExitDate"]:

    if date_column in headers:

        col = headers[date_column]

        for row in range(2, ws.max_row + 1):

            ws.cell(
                row,
                col
            ).number_format = "dd-mm-yyyy"

# ==========================================================
# ADD FILTERS
# ==========================================================

for ws in wb.worksheets:

    ws.auto_filter.ref = ws.dimensions

# ==========================================================
# SAVE
# ==========================================================

wb.save(output_file)

print()
print("========================================")
print("EXCEL WORKBOOK CREATED SUCCESSFULLY")
print("========================================")
print()
print(f"File: {output_file}")
print()
print("Sheets:")
print("1. Employee Data")
print("2. Department")
print("3. Job Roles")
print("4. Date Table")
print("5. Data Dictionary")